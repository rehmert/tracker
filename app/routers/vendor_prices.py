"""Vendor sheet comparison, including a reviewable import pipeline."""
import csv
import io
import re
from pathlib import Path
from urllib.parse import quote
from urllib.parse import urlsplit
from decimal import Decimal, ROUND_HALF_UP
from fastapi import APIRouter, Depends, File, Form, Request, UploadFile, Response
from sqlalchemy import select
from sqlalchemy.orm import Session
from ..db import get_db
from ..models import VendorPriceSheet, VendorPriceEntry, VendorShippingRule, VendorPriceAttachment
from ..templating import render
from .. import services
from ..vendor_prices import parse_lines, number, normalize_item
from ..vendor_content import parse_legacy_content, validate_content, label as content_label

router = APIRouter()
MAX_UPLOAD = 8 * 1024 * 1024

def rows_text(rows):
    buffer = io.StringIO()
    writer = csv.writer(buffer, delimiter='\t', lineterminator='\n')
    writer.writerow(('Item', 'Strength', 'Package', 'Vials', 'Total USD'))
    for item, strength, package, qty, price in rows:
        writer.writerow((item, strength, package, qty, str(price)))
    return buffer.getvalue()

def parse_review(text):
    lines = list(csv.reader(io.StringIO(text), delimiter='\t'))
    result = []
    for idx, row in enumerate(lines):
        if not row or not any(x.strip() for x in row): continue
        if idx == 0 and row[0].strip().lower() == 'item': continue
        if len(row) != 5: raise ValueError(f'Row {idx+1}: expected five tab-separated columns.')
        item, strength, package, qty, price = [x.strip() for x in row]
        if not item: raise ValueError(f'Row {idx+1}: item is required.')
        try: count = int(qty)
        except ValueError: raise ValueError(f'Row {idx+1}: vial count must be an integer.')
        value = number(price)
        if not 1 <= count <= 100000 or value is None: raise ValueError(f'Row {idx+1}: invalid count or price.')
        result.append((item[:200], strength[:80], package[:80], count, value))
    if len(result) > 500: raise ValueError('Maximum 500 items per sheet.')
    return result

def validate_product_url(value: str) -> str:
    """Only accept absolute HTTP(S) product links; empty links are optional."""
    value = value.strip()
    if not value:
        return ''
    if len(value) > 2048 or any(ch.isspace() or ord(ch) < 32 for ch in value):
        raise ValueError('Product link must be an HTTP or HTTPS URL (maximum 2048 characters).')
    try:
        parsed = urlsplit(value)
        valid = parsed.scheme.lower() in ('http', 'https') and bool(parsed.hostname)
        if parsed.username is not None or parsed.password is not None:
            valid = False
        _ = parsed.port
    except ValueError:
        valid = False
    if not valid:
        raise ValueError('Product link must be a valid HTTP or HTTPS URL.')
    return value

def shipping_for_order(rule, subtotal_cents: int) -> int:
    """Free shipping starts at the threshold (inclusive)."""
    if rule is None:
        return 0
    return 0 if subtotal_cents >= rule.free_shipping_threshold_cents else rule.flat_shipping_cents


def shipping_details(entry, rules):
    rule = rules.get(entry.sheet.vendor.strip().casefold())
    if rule is None:
        return {'configured': False, 'shipping_cents': 0, 'landed_cents': None, 'landed_per_vial_cents': None}
    shipping = shipping_for_order(rule, entry.total_price_cents)
    total = entry.total_price_cents + shipping
    return {'configured': True, 'shipping_cents': shipping, 'landed_cents': total,
            'landed_per_vial_cents': total / entry.vial_count}


# Files are archived verbatim. They are never sent to OCR or the product importer.
ARCHIVE_TYPES = {
    '.png': 'image/png', '.jpg': 'image/jpeg', '.jpeg': 'image/jpeg',
    '.webp': 'image/webp', '.csv': 'text/csv', '.tsv': 'text/tab-separated-values',
    '.xls': 'application/vnd.ms-excel',
    '.xlsx': 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
    '.ods': 'application/vnd.oasis.opendocument.spreadsheet',
}


@router.post('/vendor-prices/attachments/upload')
async def upload_attachment(request: Request, vendor: str = Form(''), title: str = Form(''),
                            upload: UploadFile = File(...), db: Session = Depends(get_db)):
    vendor = vendor.strip()
    title = title.strip()
    filename = Path((upload.filename or '').replace('\\', '/')).name
    suffix = Path(filename).suffix.lower()
    if not vendor or len(vendor) > 160 or len(title) > 160:
        return page(request, db, err='Enter a vendor (max 160 characters) and an optional title (max 160 characters).')
    if not filename or len(filename) > 255 or suffix not in ARCHIVE_TYPES:
        return page(request, db, err='Choose a PNG, JPG, WebP, CSV, TSV, XLS, XLSX, or ODS file.')
    content = await upload.read(MAX_UPLOAD + 1)
    if not content or len(content) > MAX_UPLOAD:
        return page(request, db, err='File must be nonempty and no larger than 8 MB.')
    attachment = VendorPriceAttachment(vendor=vendor, title=title, filename=filename,
                                       content_type=ARCHIVE_TYPES[suffix], content=content)
    db.add(attachment)
    db.commit()
    return services.flash_redirect('/vendor-prices', msg=f'Saved original price list for {vendor} (no products imported).')


@router.get('/vendor-prices/attachments/{attachment_id}/download')
def download_attachment(attachment_id: int, db: Session = Depends(get_db)):
    attachment = db.get(VendorPriceAttachment, attachment_id)
    if attachment is None:
        from fastapi import HTTPException
        raise HTTPException(status_code=404, detail='Price list not found.')
    # Force download; do not render untrusted uploads in the Tracker origin.
    encoded_name = quote(attachment.filename, safe='')
    return Response(content=attachment.content, media_type='application/octet-stream', headers={
        'Content-Disposition': f"attachment; filename*=UTF-8''{encoded_name}",
        'X-Content-Type-Options': 'nosniff',
        'Cache-Control': 'no-store',
    })


@router.post('/vendor-prices/attachments/{attachment_id}/delete')
def delete_attachment(attachment_id: int, db: Session = Depends(get_db)):
    attachment = db.get(VendorPriceAttachment, attachment_id)
    if attachment is None:
        return services.flash_redirect('/vendor-prices', err='Price list not found.')
    db.delete(attachment)
    db.commit()
    return services.flash_redirect('/vendor-prices', msg='Saved price list deleted.')


def page(request, db, **kwargs):
    sheets = list(db.scalars(select(VendorPriceSheet).order_by(VendorPriceSheet.id.desc())))
    entries = list(db.scalars(select(VendorPriceEntry).order_by(VendorPriceEntry.item_name, VendorPriceEntry.id)))
    groups = {}
    for entry in entries:
        key = (normalize_item(entry.item_name), entry.content_amount, entry.content_unit)
        groups.setdefault(key, []).append(entry)
    comparisons = sorted(groups.values(), key=lambda group: (group[0].item_name.lower(), (group[0].content_amount or 0), (group[0].content_unit or "")))
    shipping_rules = list(db.scalars(select(VendorShippingRule).order_by(VendorShippingRule.vendor)))
    shipping_map = {rule.vendor_key: rule for rule in shipping_rules}
    shipping_by_entry = {entry.id: shipping_details(entry, shipping_map) for entry in entries}
    vendors = sorted({s.vendor for s in sheets} | {r.vendor for r in shipping_rules}, key=str.casefold)
    attachments = list(db.scalars(select(VendorPriceAttachment).order_by(VendorPriceAttachment.id.desc())))
    return render(request, 'vendor_prices.html', sheets=sheets, attachments=attachments, comparisons=comparisons,
                  shipping_rules=shipping_rules, shipping_by_entry=shipping_by_entry,
                  shipping_vendors=vendors, **kwargs)

@router.post('/vendor-prices/shipping/save')
def save_shipping(request: Request, vendor: str = Form(''), flat_shipping: str = Form(''),
                  free_shipping_threshold: str = Form(''), db: Session = Depends(get_db)):
    vendor = vendor.strip()
    if not vendor or len(vendor) > 160:
        return page(request, db, err='Enter a vendor name (maximum 160 characters).')
    flat = number(flat_shipping)
    threshold = number(free_shipping_threshold)
    if (flat is None or threshold is None or flat < 0 or threshold <= 0
            or flat > Decimal('1000000') or threshold > Decimal('100000000')):
        return page(request, db, err='Enter a nonnegative shipping charge and a positive free-shipping threshold.')
    if flat.as_tuple().exponent < -2 or threshold.as_tuple().exponent < -2:
        return page(request, db, err='Dollar amounts must have no more than two decimal places.')
    key = vendor.casefold()
    rule = db.scalars(select(VendorShippingRule).where(VendorShippingRule.vendor_key == key)).first()
    if rule is None:
        rule = VendorShippingRule(vendor=vendor, vendor_key=key)
        db.add(rule)
    rule.vendor = vendor
    rule.flat_shipping_cents = int(flat * 100)
    rule.free_shipping_threshold_cents = int(threshold * 100)
    db.commit()
    return services.flash_redirect('/vendor-prices', msg=f'Shipping policy saved for {vendor}.')


@router.post('/vendor-prices/shipping/{rule_id}/delete')
def delete_shipping(rule_id: int, db: Session = Depends(get_db)):
    rule = db.get(VendorShippingRule, rule_id)
    if rule is None:
        return services.flash_redirect('/vendor-prices', err='Shipping policy not found.')
    db.delete(rule)
    db.commit()
    return services.flash_redirect('/vendor-prices', msg='Vendor shipping policy deleted.')


@router.get('/vendor-prices')
def index(request: Request, db: Session = Depends(get_db)):
    return page(request, db)

# Structured CSV importer: each row can have its own vendor and product link.
CSV_COLUMNS = ('vendor','item_name','content_amount','content_unit','package_type','package_price','product_url')
CSV_ALIASES = {
    'vendor': {'vendor','supplier'},
    'item_name': {'item_name','compound','item','product','name'},
    'content_amount': {'content_amount','contents_per_vial','strength','amount','mg_per_vial'},
    'content_unit': {'content_unit','unit','units'},
    'package_type': {'package_type','package','pack'},
    'package_price': {'package_price','price','total_price','price_usd'},
    'product_url': {'product_url','url','product_link','link'},
}

def parse_vendor_items_csv(data: bytes):
    text = data.decode('utf-8-sig')
    reader = csv.DictReader(io.StringIO(text))
    if not reader.fieldnames: raise ValueError('CSV must contain a header row.')
    columns = {re.sub(r'[^a-z0-9]+','_',str(h).strip().lower()).strip('_'): h for h in reader.fieldnames if h}
    mapping = {key: next((columns[a] for a in aliases if a in columns), None) for key,aliases in CSV_ALIASES.items()}
    required = ('vendor','item_name','content_amount','package_type','package_price')
    missing = [k for k in required if not mapping[k]]
    if missing: raise ValueError('Missing required CSV columns: '+', '.join(missing))
    result=[]
    for line,row in enumerate(reader,2):
        if not any((v or '').strip() for v in row.values() if isinstance(v,str)):continue
        if None in row: raise ValueError(f'Line {line}: too many columns.')
        item={key: (row.get(col) or '').strip() if col else '' for key,col in mapping.items()}
        if not item['content_unit']:item['content_unit']='mg'
        validate_vendor_csv_row(item,line)
        result.append(item)
        if len(result)>500:raise ValueError('Maximum 500 CSV items per import.')
    if not result:raise ValueError('CSV has no data rows.')
    return result

def validate_vendor_csv_row(item, line):
    if not item['vendor'] or not item['item_name']:raise ValueError(f'Line {line}: vendor and compound are required.')
    if len(item['vendor'])>160 or len(item['item_name'])>200:raise ValueError(f'Line {line}: vendor or compound name too long.')
    amount=number(item['content_amount']);price=number(item['package_price'])
    if amount is None or amount<=0 or amount>Decimal('100000000'):raise ValueError(f'Line {line}: invalid vial contents.')
    if price is None or price<0 or price>Decimal('100000000'):raise ValueError(f'Line {line}: invalid price.')
    if item['content_unit'] not in ('mg','IU'):raise ValueError(f'Line {line}: unit must be mg or IU.')
    if item['package_type'] not in ('vial','half-kit','kit'):raise ValueError(f'Line {line}: package must be vial, half-kit, or kit.')
    validate_product_url(item['product_url'])

def csv_review_text(items):
    buf=io.StringIO();writer=csv.writer(buf,delimiter='\t',lineterminator='\n')
    writer.writerow(CSV_COLUMNS)
    for item in items:writer.writerow([item[k] for k in CSV_COLUMNS])
    return buf.getvalue()

@router.post('/vendor-prices/items-csv/preview')
async def preview_items_csv(request: Request, upload: UploadFile=File(...), db: Session=Depends(get_db)):
    if not (upload.filename or '').lower().endswith('.csv'):
        return page(request,db,err='Select a CSV file.')
    data=await upload.read(MAX_UPLOAD+1)
    if len(data)>MAX_UPLOAD:return page(request,db,err='Maximum upload size is 8 MB.')
    try:items=parse_vendor_items_csv(data)
    except (ValueError,UnicodeError,csv.Error) as exc:return page(request,db,err=f'CSV import error: {str(exc)[:220]}')
    return page(request,db,csv_review=csv_review_text(items),csv_count=len(items))

@router.post('/vendor-prices/items-csv/save')
def save_items_csv(request: Request, review_text: str=Form(...), db: Session=Depends(get_db)):
    try:
        rows=list(csv.reader(io.StringIO(review_text),delimiter='\t'))
        if not rows or tuple(rows[0])!=CSV_COLUMNS:raise ValueError('Keep the seven-column header unchanged.')
        if len(rows)-1>500:raise ValueError('Maximum 500 CSV items per import.')
        items=[]
        for idx,row in enumerate(rows[1:],2):
            if not row or not any(x.strip() for x in row):continue
            if len(row)!=7:raise ValueError(f'Line {idx}: expected seven tab-separated columns.')
            item=dict(zip(CSV_COLUMNS,[x.strip() for x in row]));validate_vendor_csv_row(item,idx);items.append(item)
        if not items:raise ValueError('No items to import.')
    except (ValueError,csv.Error) as exc:return page(request,db,err=str(exc),csv_review=review_text)
    sheets={}
    for item in items:
        vendor=item['vendor']
        if vendor not in sheets:
            sheet=VendorPriceSheet(vendor=vendor,name='CSV items import')
            db.add(sheet);db.flush();sheets[vendor]=sheet
        amount=number(item['content_amount']);price=number(item['package_price'])
        strength=content_label(amount, item['content_unit'])
        package=item['package_type'];count={'vial':1,'half-kit':5,'kit':10}[package]
        db.add(VendorPriceEntry(sheet_id=sheets[vendor].id,item_name=item['item_name'],strength=strength,content_amount=float(amount),content_unit=item['content_unit'],
            package=package,vial_count=count,total_price_cents=int((price*100).quantize(Decimal('1'),rounding=ROUND_HALF_UP)),product_url=item['product_url']))
    db.commit()
    return services.flash_redirect('/vendor-prices',msg=f'Imported {len(items)} vendor offers from CSV.')

# Bulk round-trip CSV: stable entry IDs allow edits without creating duplicates.
BULK_COLUMNS = ('id', 'vendor', 'item_name', 'content_amount', 'content_unit', 'package', 'vial_count', 'total_price', 'product_url')


def bulk_export_csv(db):
    output = io.StringIO(newline='')
    writer = csv.writer(output, lineterminator='\r\n')
    writer.writerow(BULK_COLUMNS)
    entries = db.scalars(select(VendorPriceEntry).order_by(VendorPriceEntry.id)).all()
    for entry in entries:
        writer.writerow((entry.id, entry.sheet.vendor, entry.item_name, entry.content_amount if entry.content_amount is not None else "", entry.content_unit or "",
                         entry.package, entry.vial_count, f'{entry.total_price_cents / 100:.2f}',
                         entry.product_url or ''))
    return output.getvalue()


@router.get('/vendor-prices/bulk/export')
def bulk_export(db: Session = Depends(get_db)):
    return Response(content='\ufeff' + bulk_export_csv(db), media_type='text/csv; charset=utf-8', headers={
        'Content-Disposition': 'attachment; filename="tracker-vendor-prices.csv"',
        'Cache-Control': 'no-store',
    })


def parse_bulk_csv(data: bytes):
    try:
        content = data.decode('utf-8-sig')
        reader = csv.DictReader(io.StringIO(content, newline=''), strict=True)
        if reader.fieldnames is None or tuple(reader.fieldnames) != BULK_COLUMNS:
            raise ValueError('CSV header must match the exported file exactly: ' + ', '.join(BULK_COLUMNS))
        result = []
        ids = set()
        for line, row in enumerate(reader, 2):
            if None in row or any(value is None for value in row.values()):
                raise ValueError(f'Line {line}: incorrect number of columns.')
            if not any(value.strip() for value in row.values()):
                continue
            item = {key: value.strip() for key, value in row.items()}
            raw_id = item['id']
            if raw_id:
                if not raw_id.isascii() or not raw_id.isdecimal() or int(raw_id) <= 0:
                    raise ValueError(f'Line {line}: ID must be a positive integer or blank.')
                entry_id = int(raw_id)
                if entry_id in ids:
                    raise ValueError(f'Line {line}: duplicate ID {entry_id}.')
                ids.add(entry_id)
            if not item['vendor'] or len(item['vendor']) > 160:
                raise ValueError(f'Line {line}: vendor is required (maximum 160 characters).')
            if not item['item_name'] or len(item['item_name']) > 200:
                raise ValueError(f'Line {line}: item name is required (maximum 200 characters).')
            # Legacy rows may have no safely inferable per-vial content.
            # Only existing records may retain an unknown value.
            if not item['content_amount'] and not item['content_unit'] and raw_id:
                item['_amount'] = None
            else:
                amount, unit = validate_content(item['content_amount'], item['content_unit'])
                item['_amount'] = float(amount)
            if not item['package'] or len(item['package']) > 80:
                raise ValueError(f'Line {line}: invalid package description.')
            try:
                count = int(item['vial_count'])
            except ValueError:
                raise ValueError(f'Line {line}: vial count must be an integer.')
            if not 1 <= count <= 100000:
                raise ValueError(f'Line {line}: vial count must be between 1 and 100,000.')
            price = number(item['total_price'])
            if price is None or not price.is_finite() or price < 0 or price > Decimal('100000000') or price.as_tuple().exponent < -2:
                raise ValueError(f'Line {line}: price must be nonnegative USD with at most two decimal places.')
            validate_product_url(item['product_url'])
            item['_id'] = int(raw_id) if raw_id else None
            item['_count'] = count
            item['_cents'] = int(price * 100)
            result.append(item)
            if len(result) > 5000:
                raise ValueError('Maximum 5,000 rows per bulk update.')
        if not result:
            raise ValueError('CSV contains no vendor prices.')
        return result
    except (UnicodeError, csv.Error) as exc:
        raise ValueError(f'Invalid CSV: {exc}') from exc


def validate_bulk_ids(db, rows):
    ids = {row['_id'] for row in rows if row['_id'] is not None}
    existing = {e.id: e for e in db.scalars(select(VendorPriceEntry).where(VendorPriceEntry.id.in_(ids)))} if ids else {}
    missing = ids - set(existing)
    if missing:
        raise ValueError(f'Unknown vendor price ID(s): {", ".join(str(i) for i in sorted(missing)[:10])}. No changes saved.')
    for row in rows:
        if row['_amount'] is None:
            entry = existing.get(row['_id'])
            if entry is None or entry.content_amount is not None or entry.content_unit:
                raise ValueError('Blank content is allowed only for existing legacy offers with unknown content; correct the content amount and unit before saving.')
    return existing


@router.post('/vendor-prices/bulk/preview')
async def bulk_preview(request: Request, upload: UploadFile = File(...), db: Session = Depends(get_db)):
    if not (upload.filename or '').lower().endswith('.csv'):
        return page(request, db, err='Choose a CSV file exported from Tracker.')
    data = await upload.read(MAX_UPLOAD + 1)
    if len(data) > MAX_UPLOAD:
        return page(request, db, err='Maximum upload size is 8 MB.')
    try:
        rows = parse_bulk_csv(data)
        validate_bulk_ids(db, rows)
    except ValueError as exc:
        return page(request, db, err=f'Bulk update: {exc}')
    return page(request, db, bulk_review=data.decode('utf-8-sig'),
                bulk_updates=sum(row['_id'] is not None for row in rows),
                bulk_additions=sum(row['_id'] is None for row in rows))


@router.post('/vendor-prices/bulk/save')
def bulk_save(request: Request, review_text: str = Form(...), db: Session = Depends(get_db)):
    try:
        rows = parse_bulk_csv(review_text.encode('utf-8'))
        existing = validate_bulk_ids(db, rows)
        sheets = {}
        updated = added = 0
        for row in rows:
            vendor = row['vendor']
            entry = existing.get(row['_id']) if row['_id'] else None
            if entry is None or entry.sheet.vendor != vendor:
                if vendor not in sheets:
                    sheet = db.scalars(select(VendorPriceSheet).where(
                        VendorPriceSheet.vendor == vendor,
                        VendorPriceSheet.name == 'Manual entries',
                    ).order_by(VendorPriceSheet.id.desc())).first()
                    if sheet is None:
                        sheet = VendorPriceSheet(vendor=vendor, name='Manual entries')
                        db.add(sheet)
                        db.flush()
                    sheets[vendor] = sheet
                if entry is None:
                    entry = VendorPriceEntry(sheet_id=sheets[vendor].id)
                    db.add(entry)
                    added += 1
                else:
                    entry.sheet_id = sheets[vendor].id
            if row['_id']:
                updated += 1
            entry.item_name = row['item_name']
            entry.content_amount = row['_amount']
            entry.content_unit = row['content_unit']
            # Preserve original legacy strength when its content cannot be parsed.
            if row['_amount'] is not None:
                entry.strength = content_label(row['_amount'], row['content_unit'])
            entry.package = row['package']
            entry.vial_count = row['_count']
            entry.total_price_cents = row['_cents']
            entry.product_url = row['product_url']
        db.commit()
    except ValueError as exc:
        db.rollback()
        return page(request, db, err=f'Bulk update: {exc}', bulk_review=review_text)
    return services.flash_redirect('/vendor-prices', msg=f'Bulk update complete: {updated} updated, {added} added. No offers deleted.')


@router.post('/vendor-prices/preview')
async def preview(request: Request, vendor: str = Form(''), sheet_name: str = Form(''), upload: UploadFile = File(...), db: Session = Depends(get_db)):
    vendor = vendor.strip()[:160]
    if not vendor: return page(request, db, err='Vendor name is required.')
    filename = (upload.filename or '').lower()
    if not filename.endswith(('.csv', '.tsv', '.txt', '.xlsx', '.png', '.jpg', '.jpeg', '.webp')):
        return page(request, db, err='Upload CSV, TSV, XLSX, PNG, JPG or WebP.')
    data = await upload.read(MAX_UPLOAD + 1)
    if len(data) > MAX_UPLOAD: return page(request, db, err='Maximum upload size is 8 MB.')
    try:
        if filename.endswith('.xlsx'):
            from openpyxl import load_workbook
            wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
            output = io.StringIO(); writer = csv.writer(output)
            for row in wb.active.iter_rows(values_only=True):
                writer.writerow(['' if v is None else str(v) for v in row])
            wb.close(); raw = output.getvalue()
        elif filename.endswith(('.png','.jpg','.jpeg','.webp')):
            from PIL import Image
            import pytesseract
            img = Image.open(io.BytesIO(data))
            img.verify()
            img = Image.open(io.BytesIO(data))
            if img.width * img.height > 20_000_000: raise ValueError('Image too large (20 megapixels maximum).')
            raw = pytesseract.image_to_string(img)
        else:
            raw = data.decode('utf-8-sig')
        parsed = parse_lines(raw)
    except ImportError:
        return page(request, db, err='Optional parser dependency missing. Install openpyxl for XLSX, or Pillow + pytesseract + Tesseract for images.')
    except Exception as exc:
        return page(request, db, err=f'Could not parse uploaded sheet: {str(exc)[:200]}')
    return page(request, db, vendor=vendor, sheet_name=sheet_name[:160] or upload.filename, review_text=rows_text(parsed), preview_count=len(parsed), raw_text=raw[:8000])

@router.post('/vendor-prices/save')
def save(request: Request, vendor: str = Form(...), sheet_name: str = Form(''), review_text: str = Form(...), db: Session = Depends(get_db)):
    vendor = vendor.strip()[:160]
    if not vendor: return page(request, db, err='Vendor name is required.', review_text=review_text)
    try: rows = parse_review(review_text)
    except ValueError as exc: return page(request, db, err=str(exc), vendor=vendor, sheet_name=sheet_name, review_text=review_text)
    if not rows: return page(request, db, err='No valid rows to save.', vendor=vendor, sheet_name=sheet_name, review_text=review_text)
    sheet = VendorPriceSheet(vendor=vendor, name=sheet_name.strip()[:160])
    db.add(sheet); db.flush()
    for item, strength, package, qty, price in rows:
        amount, unit = parse_legacy_content(strength)
        db.add(VendorPriceEntry(sheet_id=sheet.id, item_name=item, strength=strength, content_amount=float(amount) if amount is not None else None, content_unit=unit, package=package, vial_count=qty, total_price_cents=int((price * 100).quantize(Decimal('1')))))
    db.commit()
    return services.flash_redirect('/vendor-prices', msg=f'Saved {len(rows)} offers for {vendor}.')


@router.post('/vendor-prices/manual')
def manual_add(
    request: Request,
    vendor: str = Form(''),
    item_name: str = Form(''),
    content_amount: str = Form(''),
    content_unit: str = Form('mg'),
    package_type: str = Form('vial'),
    package_price: str = Form(''),
    product_url: str = Form(''),
    db: Session = Depends(get_db),
):
    """Save one vendor offer, normalizing package price to a vial count."""
    vendor = vendor.strip()[:160]
    try:
        product_url = validate_product_url(product_url)
    except ValueError as exc:
        return page(request, db, err=str(exc))
    item_name = item_name.strip()[:200]
    amount = number(content_amount)
    price = number(package_price)
    if not vendor or not item_name:
        return page(request, db, err='Vendor and compound are required.')
    if amount is None or amount <= 0 or amount > Decimal('100000000'):
        return page(request, db, err='Vial contents must be a positive number.')
    if price is None or price > Decimal('100000000'):
        return page(request, db, err='Enter a valid nonnegative package price.')
    if content_unit not in ('mg', 'IU'):
        return page(request, db, err='Vial contents must be in mg or IU.')
    package_counts = {'vial': 1, 'half-kit': 5, 'kit': 10}
    if package_type not in package_counts:
        return page(request, db, err='Choose vial, half-kit or kit.')
    count = package_counts[package_type]
    amount_label = format(amount.normalize(), 'f')
    strength = f'{amount_label} {content_unit}/vial'
    cents = int((price * 100).quantize(Decimal('1'), rounding=ROUND_HALF_UP))
    # Reuse a manual sheet per vendor so manual entries remain editable as a group
    # via the existing sheet deletion interface; do not affect imported sheets.
    sheet = db.scalars(select(VendorPriceSheet).where(
        VendorPriceSheet.vendor == vendor,
        VendorPriceSheet.name == 'Manual entries',
    ).order_by(VendorPriceSheet.id.desc())).first()
    if sheet is None:
        sheet = VendorPriceSheet(vendor=vendor, name='Manual entries')
        db.add(sheet)
        db.flush()
    db.add(VendorPriceEntry(
        sheet_id=sheet.id, item_name=item_name, strength=strength, content_amount=float(amount), content_unit=content_unit,
        package=package_type, vial_count=count, total_price_cents=cents,
        product_url=product_url,
    ))
    db.commit()
    return services.flash_redirect('/vendor-prices', msg=f'Added {item_name} ({strength}) for {vendor}.')

@router.get('/vendor-prices/entries/{entry_id}/edit')
def edit_offer(request: Request, entry_id: int, db: Session = Depends(get_db)):
    entry = db.get(VendorPriceEntry, entry_id)
    if entry is None:
        return services.flash_redirect('/vendor-prices', err='Vendor offer not found.')
    return page(request, db, editing=entry)

@router.post('/vendor-prices/entries/{entry_id}/edit')
def update_offer(
    request: Request, entry_id: int, vendor: str = Form(''),
    item_name: str = Form(''), content_amount: str = Form(''), content_unit: str = Form('mg'),
    package: str = Form(''), vial_count: str = Form(''),
    total_price: str = Form(''), product_url: str = Form(''),
    db: Session = Depends(get_db),
):
    entry = db.get(VendorPriceEntry, entry_id)
    if entry is None:
        return services.flash_redirect('/vendor-prices', err='Vendor offer not found.')
    vendor = vendor.strip()[:160]
    item_name = item_name.strip()[:200]
    package = package.strip()[:80]
    try:
        product_url = validate_product_url(product_url)
        amount, content_unit = validate_content(content_amount, content_unit)
        count = int(vial_count)
        price = number(total_price)
        if not vendor or not item_name:
            raise ValueError('Vendor and compound are required.')
        if not package:
            raise ValueError('Package description is required.')
        if not 1 <= count <= 100000:
            raise ValueError('Vial count must be between 1 and 100,000.')
        if price is None or price < 0 or price > Decimal('100000000'):
            raise ValueError('Enter a valid nonnegative package price.')
    except (ValueError, TypeError) as exc:
        return page(request, db, editing=entry, err=str(exc))
    if vendor != entry.sheet.vendor:
        # A sheet may contain many offers: moving this entry avoids renaming
        # the vendor on unrelated offers from the same imported sheet.
        destination = db.scalars(select(VendorPriceSheet).where(
            VendorPriceSheet.vendor == vendor,
            VendorPriceSheet.name == 'Manual entries',
        ).order_by(VendorPriceSheet.id.desc())).first()
        if destination is None:
            destination = VendorPriceSheet(vendor=vendor, name='Manual entries')
            db.add(destination)
            db.flush()
        entry.sheet_id = destination.id
    entry.item_name = item_name
    entry.content_amount = float(amount)
    entry.content_unit = content_unit
    entry.strength = content_label(amount, content_unit)
    entry.package = package
    entry.vial_count = count
    entry.total_price_cents = int((price * 100).quantize(Decimal('1'), rounding=ROUND_HALF_UP))
    entry.product_url = product_url
    db.commit()
    return services.flash_redirect('/vendor-prices', msg='Vendor offer updated.')

@router.post('/vendor-prices/entries/{entry_id}/delete')
def delete_offer(entry_id: int, db: Session = Depends(get_db)):
    """Delete exactly one vendor price entry, leaving its sheet and other offers intact."""
    entry = db.get(VendorPriceEntry, entry_id)
    if entry is None:
        return services.flash_redirect('/vendor-prices', err='Vendor offer not found.')
    db.delete(entry)
    db.commit()
    return services.flash_redirect('/vendor-prices', msg='Vendor offer deleted.')

@router.post('/vendor-prices/{sheet_id}/delete')
def delete(sheet_id: int, db: Session = Depends(get_db)):
    sheet = db.get(VendorPriceSheet, sheet_id)
    if sheet:
        db.delete(sheet); db.commit()
    return services.flash_redirect('/vendor-prices', msg='Price sheet deleted.')
